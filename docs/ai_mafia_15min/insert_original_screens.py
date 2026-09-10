"""승인된 네 장의 화면 영역을 원본 JPEG의 PowerPoint 이미지로 교체한다.

ImageGen이 화면 안의 글자와 차트를 재구성하는 문제 때문에 사용자가 직접 삽입을
승인했다. 원본 바이트를 수정하지 않고 PowerPoint의 자르기 속성만 사용하며,
생성된 화면이 여백에 남지 않도록 교체 영역 전체를 먼저 가린다.
"""

from pathlib import Path
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement


BASE = Path(__file__).resolve().parent
PPTX = BASE / "ai_mafia_15min.pptx"
REFERENCE_WIDTH = 1672
REFERENCE_HEIGHT = 941

# 각 영역은 생성된 배경의 스크린샷 내부만 포함한다. 제목과 외부 설명은 유지한다.
PLACEMENTS = {
    4: [
        ("home.jpg", (0, 40, 1280, 350), (51, 344, 925, 298)),
        ("role_reveal.jpg", (325, 270, 625, 405), (1029, 320, 592, 371)),
    ],
    8: [
        ("game_discussion.jpg", (0, 0, 1280, 720), (48, 235, 728, 570)),
        ("game_chat.jpg", (352, 210, 865, 370), (813, 243, 810, 569)),
    ],
    9: [
        ("custom_role_setup.jpg", (75, 85, 1115, 295), (65, 240, 792, 250)),
        ("custom_role.jpg", (75, 398, 490, 83), (64, 573, 793, 240)),
    ],
    11: [
        ("admin_overview.jpg", (0, 320, 1280, 390), (45, 118, 928, 329)),
        ("admin_speech_analytics.jpg", (0, 150, 1280, 560), (45, 472, 928, 399)),
    ],
}


def add_original(prs, slide, filename, crop, region):
    """원본 종횡비를 유지해 영역에 맞추고, 화면에서 보일 범위만 지정한다."""
    x, y, w, h = region
    left = round(x / REFERENCE_WIDTH * prs.slide_width)
    top = round(y / REFERENCE_HEIGHT * prs.slide_height)
    width = round(w / REFERENCE_WIDTH * prs.slide_width)
    height = round(h / REFERENCE_HEIGHT * prs.slide_height)
    mask = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    mask.name = f"원본 화면 교체 배경: {filename}"
    mask.fill.solid()
    mask.fill.fore_color.rgb = RGBColor.from_string("F7F8FC")
    mask.line.fill.background()
    # 템플릿의 기본 그림자가 화면 테두리처럼 보이지 않도록 효과 상속을 끈다.
    mask._element.spPr.append(OxmlElement("a:effectLst"))
    sx, sy, sw, sh = crop
    picture_width = min(width, round(height * sw / sh))
    picture_height = round(picture_width * sh / sw)
    picture = slide.shapes.add_picture(
        str(BASE / "assets" / "screenshots" / filename),
        left + (width - picture_width) // 2,
        top + (height - picture_height) // 2,
        width=picture_width,
        height=picture_height,
    )
    picture.name = f"원본 캡처: {filename}"
    picture._element.nvPicPr.cNvPr.set("descr", f"실제 실행 화면 원본: {filename}")
    iw, ih = picture.image.size
    assert 0 <= sx < sx + sw <= iw and 0 <= sy < sy + sh <= ih
    picture.crop_left = sx / iw
    picture.crop_right = (iw - sx - sw) / iw
    picture.crop_top = sy / ih
    picture.crop_bottom = (ih - sy - sh) / ih


def main():
    """사용자 승인과 기본 조립 상태를 확인한 다음 화면 여덟 개를 삽입한다."""
    spec = json.loads((BASE / "deck_spec.json").read_text())
    assert spec["source_screenshot_exception"]["approved_by_user"]
    prs = Presentation(PPTX)
    assert len(prs.slides) == 16
    assert all(len(slide.shapes) == 1 for slide in prs.slides), "기본 조립 파일에만 실행하세요."
    for number, items in PLACEMENTS.items():
        for filename, crop, region in items:
            add_original(prs, prs.slides[number - 1], filename, crop, region)
    prs.core_properties.title = "AI MAFIA — 15분 프로젝트 발표"
    prs.core_properties.subject = "AI Agent·MCP·분리 실행 구조와 실제 게임·관리자 화면"
    prs.core_properties.author = "Team 4"
    prs.save(PPTX)
    print("원본 스크린샷 삽입 완료: 4·8·9·11번, 총 8개")


if __name__ == "__main__":
    main()
